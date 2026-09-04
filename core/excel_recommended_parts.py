"""Importación Excel de partes recomendadas por subcategoría."""

from __future__ import annotations

from typing import Any

from openpyxl import load_workbook

from .models import SubcategoryProduct, SubcategoryRecommendedPart

REQUIRED_HEADERS = {"subcategory_id", "name"}
OPTIONAL_HEADERS = {"description", "sort_order", "is_active"}


def _cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_bool(value, default: bool = True) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    s = str(value).strip().lower()
    if s in ("1", "true", "yes", "si", "sí", "y", "activo"):
        return True
    if s in ("0", "false", "no", "n", "inactivo"):
        return False
    return default


def import_recommended_parts_from_excel(file_obj) -> dict[str, Any]:
    """
    Columnas esperadas (fila 1 = headers):
      subcategory_id (requerido)
      name (requerido)
      description (opcional)
      sort_order (opcional, int)
      is_active (opcional: true/false, 1/0, si/no)

    Upsert por (subcategory_id, name).
    """
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return {"created": 0, "updated": 0, "errors": [{"row": 1, "error": "Archivo vacío"}]}

    headers = [_cell_str(h).lower() for h in header_row]
    header_set = {h for h in headers if h}
    missing = REQUIRED_HEADERS - header_set
    if missing:
        return {
            "created": 0,
            "updated": 0,
            "errors": [
                {
                    "row": 1,
                    "error": f"Faltan columnas: {', '.join(sorted(missing))}",
                }
            ],
        }

    col = {name: idx for idx, name in enumerate(headers)}
    created = 0
    updated = 0
    errors: list[dict] = []

    for row_idx, row in enumerate(rows, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        try:
            subcategory_raw = row[col["subcategory_id"]]
            name = _cell_str(row[col["name"]])
            if subcategory_raw is None or name == "":
                errors.append({"row": row_idx, "error": "subcategory_id y name son obligatorios"})
                continue
            subcategory_id = int(subcategory_raw)
            if not SubcategoryProduct.objects.filter(pk=subcategory_id).exists():
                errors.append(
                    {"row": row_idx, "error": f"subcategory_id={subcategory_id} no existe"}
                )
                continue

            description = ""
            if "description" in col and col["description"] < len(row):
                description = _cell_str(row[col["description"]])

            sort_order = 0
            if "sort_order" in col and col["sort_order"] < len(row) and row[col["sort_order"]] not in (
                None,
                "",
            ):
                sort_order = int(row[col["sort_order"]])

            is_active = True
            if "is_active" in col and col["is_active"] < len(row):
                is_active = _parse_bool(row[col["is_active"]], default=True)

            obj, was_created = SubcategoryRecommendedPart.objects.update_or_create(
                subcategory_id=subcategory_id,
                name=name,
                defaults={
                    "description": description,
                    "sort_order": sort_order,
                    "is_active": is_active,
                },
            )
            if was_created:
                created += 1
            else:
                updated += 1
        except Exception as exc:  # noqa: BLE001 — report per-row errors
            errors.append({"row": row_idx, "error": str(exc)})

    return {"created": created, "updated": updated, "errors": errors}
